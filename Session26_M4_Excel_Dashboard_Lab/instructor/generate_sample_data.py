"""
Generates the SIMULATED sample files for the Session 26 / M4 lab (Kesar Kitchen Foods).
Run from the lab folder:  python instructor/generate_sample_data.py
Output: 01_sample_data/*.xlsx and *.csv  (fixed random seed -> same files every time)

The files are deliberately a little messy (like real exports) so that students must
ask the AI to handle: title rows, TOTAL rows, blank rows, duplicate rows, amounts stored
as text ("Rs 1,250"), dates stored as text, inconsistent city spellings.
"""
import csv
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

random.seed(26)
OUT = Path(__file__).resolve().parent.parent / "01_sample_data"
OUT.mkdir(exist_ok=True)

HEAD_FONT = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="7A3E00")

PRODUCTS = [
    # name, category, MRP
    ("Bhavnagari Gathiya 400g", "Namkeen", 120),
    ("Ratlami Sev 400g", "Namkeen", 110),
    ("Khatta Meetha Mix 400g", "Namkeen", 130),
    ("Kesar Peda 500g", "Sweets", 350),
    ("Kaju Katli 500g", "Sweets", 520),
    ("Mohanthal 500g", "Sweets", 300),
    ("Instant Dhokla Mix 500g", "Ready-to-Eat", 95),
    ("Khandvi Mix 500g", "Ready-to-Eat", 105),
    ("Thepla Pack 10pcs", "Ready-to-Eat", 80),
    ("Kesar Garam Masala 100g", "Spices", 75),
    ("Chhundo 400g", "Spices", 160),
    ("Methia Masala 200g", "Spices", 90),
    ("Kesar Badam Milk Mix 200g", "Beverages", 180),
    ("Masala Chai Premix 500g", "Beverages", 220),
    ("Aam Panna 750ml", "Beverages", 140),
]
CITIES = ["Vadodara", "Ahmedabad", "Surat", "Rajkot"]
CITY_W = [35, 35, 18, 12]
CITY_VARIANTS = {
    "Vadodara": ["vadodara", "Baroda", "VADODARA", "Vadodara "],
    "Ahmedabad": ["AHMEDABAD", "Ahmd", "ahmedabad", "Amdavad"],
    "Surat": ["surat", "SURAT"],
    "Rajkot": ["rajkot", "Rajkot "],
}
CHANNELS = ["Retail Store", "Distributor", "Amazon", "Own Website", "Quick Commerce"]
REPS = {
    "Vadodara": ["Hetal Shah", "Manan Desai"],
    "Ahmedabad": ["Ritu Joshi", "Kunal Mehta"],
    "Surat": ["Nirav Patel"],
    "Rajkot": ["Dhara Vyas"],
}
MONTHS = [4, 5, 6, 7, 8, 9]
MONTH_BOOST = {4: 0.85, 5: 0.9, 6: 0.95, 7: 1.0, 8: 1.35, 9: 1.15}  # Raksha Bandhan in Aug


def rand_date(month):
    start = date(2026, month, 1)
    end = date(2026, month + 1, 1) if month < 12 else date(2027, 1, 1)
    return start + timedelta(days=random.randrange((end - start).days))


def style_header(ws, row=1):
    for c in ws[row]:
        if c.value is not None:
            c.font = HEAD_FONT
            c.fill = HEAD_FILL


def autosize(ws):
    for col in ws.columns:
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 45)


# ---------------------------------------------------------------- 1. SALES
def make_sales():
    rows = []
    oid = 50001
    for m in MONTHS:
        n = int(230 * MONTH_BOOST[m])
        for _ in range(n):
            city = random.choices(CITIES, CITY_W)[0]
            # Quick Commerce grows every month (a story students should find)
            qc_w = 4 + (m - 4) * 3
            ch = random.choices(CHANNELS, [34, 18, 16, 10, qc_w])[0]
            prod, cat, price = random.choice(PRODUCTS)
            if m == 8 and random.random() < 0.3:
                prod, cat, price = random.choice(PRODUCTS[3:6])  # festive sweets
            if ch == "Distributor":
                qty, disc = random.randint(20, 120), random.choice([12, 15, 18])
            elif ch == "Retail Store":
                qty, disc = random.randint(5, 40), random.choice([5, 8, 10])
            elif ch == "Quick Commerce":
                qty, disc = random.randint(1, 4), random.choice([0, 0, 5])
            else:
                qty, disc = random.randint(1, 6), random.choice([0, 5, 10, 15])
            if city == "Surat" and ch in ("Retail Store", "Distributor"):
                qty = max(1, int(qty * 0.7))  # Surat under-performs
            net = round(qty * price * (1 - disc / 100))
            rep = random.choice(REPS[city]) if ch in ("Retail Store", "Distributor") else "Online Team"
            pay = (random.choice(["Credit 30 days", "NEFT"]) if ch == "Distributor"
                   else random.choice(["UPI", "UPI", "Cash", "Card"]) if ch == "Retail Store"
                   else random.choice(["UPI", "Card", "COD"]))
            rows.append([f"KK-{oid}", rand_date(m), city, ch, prod, cat, qty, price, disc, net, rep, pay])
            oid += 1
    rows.sort(key=lambda r: r[1])
    for i, r in enumerate(rows):
        r[0] = f"KK-{50001 + i}"

    clean_rows = [list(r) for r in rows]  # truth before messing up (used by answer key)

    # --- make it messy
    for r in rows:
        if random.random() < 0.08:
            r[2] = random.choice(CITY_VARIANTS[r[2]])
        if random.random() < 0.05:
            r[1] = r[1].strftime(random.choice(["%d-%m-%Y", "%d/%m/%Y"]))
        if random.random() < 0.04:
            r[9] = f"Rs {r[9]:,}"
    dup_idx = random.sample(range(len(rows)), 7)
    for i in sorted(dup_idx, reverse=True):
        rows.insert(i + 1, list(rows[i]))  # exact duplicate right after
    for i in random.sample(range(50, len(rows) - 50), 3):
        rows.insert(i, [None] * 12)  # blank rows

    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(["Order_ID", "Order_Date", "City", "Channel", "Product", "Category", "Qty",
               "Unit_Price", "Discount_Pct", "Net_Amount", "Sales_Rep", "Payment_Mode"])
    for r in rows:
        ws.append(r)
    for c in ws["B"][1:]:
        if isinstance(c.value, date):
            c.number_format = "DD-MMM-YYYY"
    style_header(ws)
    autosize(ws)
    ws.freeze_panes = "A2"

    # Targets sheet: based on true actuals, Surat set too high
    actual = {}
    for r in clean_rows:
        k = (r[1].month, r[2])
        actual[k] = actual.get(k, 0) + r[9]
    t = wb.create_sheet("Monthly_Targets")
    t.append(["Month", "City", "Target_Amount"])
    for m in MONTHS:
        for city in CITIES:
            f = random.uniform(1.25, 1.4) if city == "Surat" else random.uniform(0.9, 1.08)
            t.append([date(2026, m, 1), city, int(round(actual[(m, city)] * f, -3))])
    for c in t["A"][1:]:
        c.number_format = "MMM-YYYY"
    style_header(t)
    autosize(t)

    wb.save(OUT / "Sales_Apr-Sep-2026.xlsx")
    return clean_rows


# ---------------------------------------------------------------- 2. INVENTORY
def make_inventory():
    wb = Workbook()
    ws = wb.active
    ws.title = "Stock_Sep-2026"
    ws.append(["Kesar Kitchen Foods - Warehouse Stock Report as of 30-Sep-2026 (units)"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["SKU", "Product", "Category", "Warehouse", "Opening_Stock", "Received", "Sold",
               "Closing_Stock", "Reorder_Level", "Unit_Cost", "Expiry_Date"])
    style_header(ws, 3)
    whs = ["Vadodara WH", "Ahmedabad WH", "Surat WH"]
    n = 0
    mismatch_rows = random.sample(range(45), 2)
    for pi, (prod, cat, price) in enumerate(PRODUCTS):
        for w in whs:
            sku = f"KK{pi + 1:02d}-{w[:3].upper()}"
            opening = random.randint(150, 900)
            received = random.choice([0, 200, 300, 500, 800])
            sold = random.randint(100, opening + received - 10)
            if random.random() < 0.18:  # some SKUs run low
                sold = opening + received - random.randint(5, 40)
            closing = opening + received - sold
            if n in mismatch_rows:
                closing += random.choice([50, -30])  # data-entry error to catch with a formula
            reorder = random.choice([80, 100, 120, 150])
            cost = round(price * random.uniform(0.55, 0.65))
            shelf = 25 if cat == "Sweets" else 70 if cat == "Ready-to-Eat" else 160
            expiry = date(2026, 9, 30) + timedelta(days=random.randint(shelf // 3, shelf * 2))
            ws.append([sku, prod, cat, w, opening, received, sold, closing, reorder, cost, expiry])
            n += 1
    # exported TOTAL row (static numbers, like a real ERP export)
    tot = [sum(ws.cell(r, c).value for r in range(4, ws.max_row + 1)) for c in range(5, 9)]
    ws.append(["TOTAL", None, None, None, *tot, None, None, None])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for c in ws["K"][3:]:
        c.number_format = "DD-MMM-YYYY"
    autosize(ws)
    ws.column_dimensions["A"].width = 14
    wb.save(OUT / "Inventory_Stock_Sep-2026.xlsx")


# ---------------------------------------------------------------- 3. MARKETING
def make_marketing():
    platforms = {
        # platform: (spend range, CPM-ish impressions per Rs, CTR, lead rate, order rate, AOV)
        "Instagram": ((15000, 25000), 55, 0.012, 0.08, 0.10, 420),
        "Facebook": ((8000, 15000), 45, 0.010, 0.08, 0.10, 400),
        "Google Ads": ((12000, 20000), 12, 0.045, 0.10, 0.15, 520),
        "YouTube": ((10000, 18000), 70, 0.004, 0.06, 0.10, 380),
        "WhatsApp": ((3000, 6000), 6, 0.20, 0.12, 0.12, 560),
        "Print / Hoardings": ((12000, 25000), 30, 0.001, 0.25, 0.15, 350),
    }
    names = {4: "Summer Aam Panna", 5: "Vacation Snack Box", 6: "Monsoon Chai", 7: "Shravan Fasting Range",
             8: "Raksha Bandhan Sweets", 9: "Navratri Pre-Orders"}
    wb = Workbook()
    ws = wb.active
    ws.title = "Campaigns"
    ws.append(["Campaign_ID", "Campaign_Name", "Platform", "Month", "Spend", "Impressions", "Clicks",
               "Leads", "Orders", "Revenue"])
    cid = 1
    for m in MONTHS:
        for p, (sr, ipr, ctr, lr, orr, aov) in platforms.items():
            spend = round(random.uniform(*sr) * (1.3 if m in (8, 9) else 1), -2)
            imp = int(spend * ipr * random.uniform(0.85, 1.15))
            clicks = int(imp * ctr * random.uniform(0.8, 1.2))
            leads = int(clicks * lr * random.uniform(0.8, 1.2))
            orders = int(leads * orr * random.uniform(0.8, 1.2))
            rev = int(orders * aov * random.uniform(0.9, 1.1))
            spend_val = f"Rs {int(spend):,}" if random.random() < 0.08 else int(spend)
            ws.append([f"MK-{cid:03d}", f"{names[m]} - {p}", p, date(2026, m, 1), spend_val, imp, clicks,
                       leads, orders, rev])
            cid += 1
    for c in ws["D"][1:]:
        c.number_format = "MMM-YYYY"
    style_header(ws)
    autosize(ws)
    wb.save(OUT / "Marketing_Campaigns_H1.xlsx")


# ---------------------------------------------------------------- 4. FEEDBACK (CSV)
def make_feedback():
    issues = ["Taste", "Packaging", "Delivery delay", "Price", "Quantity / weight", "None"]
    comments = {
        "Taste": ["Too salty this time", "Not as fresh as before", "Taste was different from last pack"],
        "Packaging": ["Pack was torn", "Zip lock not working", "Box was crushed"],
        "Delivery delay": ["Order came 4 days late", "Delivery partner did not call", "Late delivery"],
        "Price": ["Costly compared to local brand", "Price increased again"],
        "Quantity / weight": ["Packet felt less than 500g", "Less pieces in box"],
        "None": ["Loved it", "Very tasty, will reorder", "Good quality", "Authentic Gujarati taste", ""],
    }
    rows = []
    for i in range(430):
        m = random.choice(MONTHS)
        d = rand_date(m)
        city = random.choices(CITIES, CITY_W)[0]
        ch = random.choices(CHANNELS, [30, 5, 25, 20, 20])[0]
        prod = random.choice(PRODUCTS)[0]
        w = [10, 8, 6, 6, 5, 65]
        if ch == "Own Website":
            w[2] = 30  # website deliveries are slow
        if ch == "Amazon":
            w[1] = 22  # packaging damage
        if city == "Surat":
            w[2] += 10
        issue = random.choices(issues, w)[0]
        if issue == "None":
            rating = random.choices([3, 4, 5], [10, 40, 50])[0]
            nps = random.choices(range(6, 11), [5, 10, 25, 30, 30])[0]
        else:
            rating = random.choices([1, 2, 3, 4], [25, 35, 30, 10])[0]
            nps = random.choices(range(0, 9), [5, 8, 10, 12, 15, 15, 15, 10, 10])[0]
        rows.append([None, d.strftime("%d-%m-%Y"), city, ch, prod,
                     "" if random.random() < 0.03 else rating, nps, issue,
                     random.choice(comments[issue])])
    rows.sort(key=lambda r: (r[1][6:], r[1][3:5], r[1][:2]))
    for i, r in enumerate(rows):
        r[0] = f"FB-{i + 1:04d}"
    with open(OUT / "Customer_Feedback_H1.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["Feedback_ID", "Date", "City", "Channel", "Product", "Rating", "NPS_Score",
                    "Issue_Type", "Comment"])
        w.writerows(rows)


if __name__ == "__main__":
    make_sales()
    make_inventory()
    make_marketing()
    make_feedback()
    print("Sample files written to", OUT)
